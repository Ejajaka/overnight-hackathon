"""Generate docs/Architecture_and_Proposed_Solution.docx.

Run: python -m scripts.make_docs
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "docs" / "Architecture_and_Proposed_Solution.docx"

GREEN = RGBColor(0x07, 0x5E, 0x54)
DARK = RGBColor(0x1A, 0x1A, 0x1A)


def code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    doc.add_paragraph()
    return t


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.color.rgb = GREEN
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.color.rgb = GREEN
    return p


def build():
    doc = Document()
    doc.core_properties.title = "Source-Truth Listings - Architecture and Proposed Solution"
    doc.core_properties.author = "Source-Truth Listings team"

    title = doc.add_heading("Source-Truth Listings", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Architecture and Proposed Solution")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(15)
        run.font.color.rgb = DARK
    tagline = doc.add_paragraph(
        "Turning an artisan's spoken knowledge into a trustworthy, buyer-ready listing in minutes - "
        "where every claim is traceable to the maker and nothing is invented."
    )
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # 1. Executive summary
    h1(doc, "1. Executive Summary")
    doc.add_paragraph(
        "Source-Truth Listings is a WhatsApp-first assistant that interviews a craftsperson in plain "
        "conversation, captures what they actually say as confirmed facts, uses a generative AI model to "
        "draft a polished marketplace listing, and then runs a claim-level 'source-truth guard' that blocks "
        "or repairs any statement the maker never confirmed. The result is a buyer-ready listing whose every "
        "material, care, production, and cultural claim is traceable to the maker. Buyers receive a shareable "
        "web page that answers their real questions (is the photo the exact item, is it machine washable, why "
        "does it take longer, what will vary) without fabricated claims."
    )
    doc.add_paragraph(
        "The team's core technical contribution is the guard: a hybrid of an LLM claim auditor and a "
        "deterministic, default-deny enforcement layer. It is the difference between 'an AI wrote a nice "
        "description' and 'an AI wrote a description that is safe to publish for a traditional craft community'."
    )

    # 2. Background
    h1(doc, "2. Background and Problem Statement")
    h2(doc, "2.1 The scenario")
    doc.add_paragraph(
        "A craftsperson creates handwoven products using techniques learned within the family. Their online "
        "listings describe little more than colour and price. Buyers compare the work with mass-produced goods "
        "and repeatedly ask why delivery takes longer."
    )
    h2(doc, "2.2 The information and trust gap")
    bullets(
        doc,
        [
            "Buyers cannot tell whether the photograph is the exact item they will receive.",
            "Buyers do not know how the product should be cared for (for example, whether it is machine washable).",
            "Buyers do not understand why handmade products take longer to produce.",
            "Buyers do not know what natural variations to expect between pieces.",
            "As a result: repeated questions, abandoned purchases, returns, and unfair comparisons with mass production.",
        ],
    )
    h2(doc, "2.3 Why existing AI writing tools fail here")
    doc.add_paragraph(
        "Generic AI copywriters can produce polished text, but an LLM will happily invent care instructions, "
        "material properties, or cultural meaning to fill gaps. For traditional crafts an invented cultural "
        "explanation is not merely inaccurate: it can misrepresent a community's heritage."
    )
    h2(doc, "2.4 Focused problem statement (finalized)")
    p = doc.add_paragraph()
    run = p.add_run(
        "How might we convert an artisan's spoken knowledge into a trustworthy, buyer-ready listing in "
        "minutes, where every claim is traceable to the maker and nothing is invented?"
    )
    run.bold = True
    run.font.size = Pt(12)
    h2(doc, "2.5 Constraints from the scenario")
    table(
        doc,
        ["Constraint", "Implication for the design"],
        [
            ("Basic phone, limited English, little time", "Conversational, low-typing, WhatsApp-first; quick-reply buttons; optional voice notes."),
            ("Materials and patterns vary naturally", "Listings must describe variation instead of implying mass-production consistency."),
            ("Two handmade pieces are rarely identical", "Per-piece 'what will vary' and 'the exact piece' framing."),
            ("Patterns carry cultural meaning", "Cultural claims may only be restated when the maker confirms them; never invented."),
            ("Community concern about misrepresentation", "Default-deny guard with explicit, respectful fallbacks."),
        ],
    )

    # 3. Proposed solution
    h1(doc, "3. Proposed Solution")
    doc.add_paragraph(
        "Source-Truth Listings is a WhatsApp-style assistant plus a shareable buyer page. The maker chats with "
        "a bot exactly as they would with a person; the assistant does the administrative work in the background "
        "and produces a listing that is safe to publish."
    )
    h2(doc, "3.1 How it addresses each buyer question")
    table(
        doc,
        ["Buyer question", "How the solution answers it (only from confirmed facts)"],
        [
            ("Is the photograph the exact item?", "A per-piece note is published only if the maker confirmed it; otherwise it is omitted or flagged."),
            ("Is this machine washable?", "Care is stated only from a closed set of care facts the maker confirmed; otherwise 'Care for this piece is not yet confirmed - ask the maker.'"),
            ("Why does delivery take longer?", "Production time is restated from the maker's own description of the process."),
            ("What natural variations will I see?", "Variation facts restated from the maker, framed as part of the handmade value."),
            ("Does the pattern have meaning?", "Cultural meaning is restated only if the maker confirmed it; invented heritage claims are blocked."),
        ],
    )
    h2(doc, "3.2 The differentiator")
    doc.add_paragraph(
        "The differentiator is not the writing; it is provenance. Every published sentence maps to a confirmed "
        "maker fact. The system can prove why a sentence is in the listing, and can show which sentences it "
        "refused to publish."
    )

    # 4. Architecture
    h1(doc, "4. System Architecture")
    h2(doc, "4.1 High-level architecture")
    code(
        doc,
        "                         MAKER (WhatsApp)                         BUYER (browser)\n"
        "                                |                                      |\n"
        "                     voice note / photo / text                 shareable buyer link\n"
        "                                |                                      |\n"
        "                                v                                      v\n"
        "                 +----------------------------+          +---------------------------+\n"
        "                 |  Twilio WhatsApp (trial)   |          |  Buyer page /buyer/{id}   |\n"
        "                 +----------------------------+          +---------------------------+\n"
        "                                |  POST /webhook/twilio                 ^\n"
        "                                v                                       |  GET /api/listing/{id}\n"
        "      +---------------------------------------------------------------------------+\n"
        "      |                       FastAPI application (app.py)                       |\n"
        "      |                                                                           |\n"
        "      |   Transport        Agent orchestrator        Session store (in-memory)    |\n"
        "      |   (Twilio/TwiML) -> interview -> ledger ->  draft -> GUARD -> reply      |\n"
        "      |            |                         |            |            |         |\n"
        "      |            |                         v            v            v         |\n"
        "      |            |                    Claim Ledger   Draft module  Guard       |\n"
        "      |            |                    (confirmed     (LLM writer)  (claim       |\n"
        "      |            |                     facts)                      audit +     |\n"
        "      |            |                                                 repair)     |\n"
        "      |            v                                                              |\n"
        "      |   OpenAI API (chat + Whisper)      Media store (photo bytes -> /media/{id})|\n"
        "      +---------------------------------------------------------------------------+\n"
        "                                ^\n"
        "                                |  HTTPS tunnel (cloudflared / ngrok)\n"
        "                                |  https://<random>.trycloudflare.com/webhook/twilio",
    )
    h2(doc, "4.2 Component responsibilities")
    table(
        doc,
        ["Component", "Responsibility"],
        [
            ("Transport / webhook", "Receives Twilio form posts, downloads media (photo/voice), converts to an internal message, returns TwiML replies."),
            ("Agent orchestrator", "Owns the conversation state machine: idle -> interview -> confirm -> review -> publish."),
            ("Interview module", "Guided questions (identity, material, care, process, variation, photo, cultural) with quick replies."),
            ("Claim Ledger", "Stores only confirmed facts: {id, type, text, source_turn, confirmed}. The single source of truth."),
            ("Draft module", "Uses an LLM to write a complete, persuasive listing from the confirmed facts."),
            ("Guard (core)", "Splits the draft into atomic claims, classifies each against the ledger, and repairs or blocks unsafe claims."),
            ("Session store", "In-memory per-conversation state; Twilio sessions are keyed by phone number."),
            ("Buyer page", "Read-only listing with a provenance panel showing the confirmed facts used."),
            ("Simulator UI", "A browser WhatsApp clone so the whole flow can be tested without a phone."),
        ],
    )
    h2(doc, "4.3 Request/response flow (maker)")
    numbered(
        doc,
        [
            "Maker messages the WhatsApp number. Twilio posts the inbound message to /webhook/twilio.",
            "Transport extracts text; if media is present it downloads the image (stored) or transcribes the audio (Whisper).",
            "Agent classifies the message into the current conversation stage and asks the next question.",
            "Each answer is extracted into facts and echoed back for a one-tap confirm; only confirmed facts enter the Ledger.",
            "After the last question, the Draft module writes a listing and the Guard audits it claim by claim.",
            "The Guard returns a safe listing plus a list of blocked claims; the agent sends both to the maker.",
            "Maker replies Publish; the agent sends the buyer link back over WhatsApp.",
        ],
    )
    h2(doc, "4.4 Deployment view")
    bullets(
        doc,
        [
            "Single process: `python app.py` runs the FastAPI app and, with `--live`, also starts the public tunnel.",
            "Local: http://127.0.0.1:8000 (simulator + buyer pages).",
            "Public: an HTTPS tunnel URL (cloudflared trycloudflare or ngrok) used as the Twilio webhook target.",
            "Configuration through environment variables / .env (secrets never committed).",
        ],
    )

    # 5. Guard
    h1(doc, "5. The Source-Truth Guard (Core Contribution)")
    doc.add_paragraph(
        "The guard is what makes the listing publishable. It is deliberately conservative: when in doubt, a claim "
        "is blocked, not softened."
    )
    h2(doc, "5.1 Pipeline")
    numbered(
        doc,
        [
            "Split: break every listing field into atomic claims (smallest meaningful statements).",
            "Classify: label each claim's type (material, care, cultural, process, variation, photo, general, neutral).",
            "Judge: an LLM auditor marks status as supported, unsupported, or cultural_unverified against the ledger.",
            "Enforce: a deterministic layer overrides the auditor for the sensitive classes. Care claims require a confirmed care fact; cultural claims require a confirmed cultural fact. This is the safety net that does not trust the LLM.",
            "Repair: supported/neutral claims are kept; unsupported claims are dropped or replaced with a closed, respectful fallback.",
            "Reassemble: rebuild each field from the surviving claims and preserve the title.",
        ],
    )
    h2(doc, "5.2 Statuses and behaviour")
    table(
        doc,
        ["Status", "Meaning", "Published as"],
        [
            ("supported", "Ledger explicitly states it; evidence recorded.", "Kept, with the matching fact as provenance."),
            ("unsupported", "Invented or unverifiable (for example an invented care instruction).", "Dropped, or replaced with a safe fallback."),
            ("cultural_unverified", "A cultural/heritage claim with no confirmed source.", "Never described; replaced with 'holds significance... details not documented'."),
            ("neutral", "Non-factual pleasantries.", "Kept."),
        ],
    )
    h2(doc, "5.3 Closed fallbacks (examples)")
    code(
        doc,
        "Care (unconfirmed)  -> \"Care for this piece is not yet confirmed - ask the maker.\"\n"
        "Cultural (unverified)-> \"This pattern holds significance in the maker's community; its detailed\n"
        "                         meaning is not documented here.\"",
    )
    h2(doc, "5.4 Why a hybrid guard, not just an LLM")
    bullets(
        doc,
        [
            "LLMs are good at understanding and splitting text, but unreliable as a final authority on truth.",
            "The deterministic layer guarantees care and cultural safety even if the auditor makes a mistake.",
            "Closed fallbacks mean an uncertain answer degrades gracefully instead of becoming a false claim.",
            "Provenance (evidence spans) makes the result auditable and explainable to the maker and to buyers.",
        ],
    )

    # 6. Data model
    h1(doc, "6. Data Model")
    h2(doc, "6.1 Claim Ledger fact")
    code(
        doc,
        "{ id: int, type: 'material|care|process|variation|photo|cultural|identity|provenance',\n"
        "  text: str, source_turn: int, confirmed: true }",
    )
    h2(doc, "6.2 Audited claim")
    code(
        doc,
        "{ field: str, claim: str, type: str,\n"
        "  status: 'supported|unsupported|cultural_unverified',\n"
        "  evidence: str, replacement: str }",
    )
    h2(doc, "6.3 Published listing fields")
    code(
        doc,
        "title, story, materials, care, production, variations,\n"
        "cultural_note, photo_note, buyer_faq",
    )
    h2(doc, "6.4 Session")
    code(
        doc,
        "{ id, ledger, stage, q_index, pending_facts, audit, published, photo, created }",
    )

    # 7. Why we chose
    h1(doc, "7. Technology Choices and Rationale (Why we chose what we chose)")
    h2(doc, "7.1 Generative AI / LLM (OpenAI) - required technology")
    bullets(
        doc,
        [
            "Why LLM: the task is unstructured natural language (spoken knowledge) -> structured facts -> persuasive copy. This is exactly what LLMs are good at.",
            "Why OpenAI specifically: single provider for both chat (JSON structured output) and speech-to-text (Whisper), a mature Python SDK, and reliable JSON mode for the guard's structured claims.",
            "Why structured JSON output: the guard needs machine-checkable claims, not free prose.",
            "Why a mock fallback: the whole product must demo offline and never break during a live presentation; if no key is set or the API fails, a deterministic mock runs.",
        ],
    )
    h2(doc, "7.2 Claim-level guard instead of a single 'be truthful' prompt")
    bullets(
        doc,
        [
            "A polite instruction ('do not hallucinate') is not a guarantee. We wanted a verifiable mechanism.",
            "Atomic claims let us audit, block, and show provenance at sentence level.",
            "Deterministic enforcement for care/cultural removes reliance on the model's goodwill.",
        ],
    )
    h2(doc, "7.3 WhatsApp via Twilio Sandbox")
    table(
        doc,
        ["Option", "Why / why not"],
        [
            ("Twilio WhatsApp Sandbox / trial sender (chosen)", "Live in minutes, no Meta business verification, maker keeps ownership of the Twilio account, supports inbound text, images and voice."),
            ("Meta WhatsApp Cloud API (deferred)", "Production-grade and own number, but needs a Meta Business account and app review (often days). Kept as a future provider behind the same transport layer."),
            ("Custom app only (rejected)", "The artisan has a basic phone and little time; a native app adds friction and installation."),
        ],
    )
    h2(doc, "7.4 Public tunnel: cloudflared first, ngrok second")
    bullets(
        doc,
        [
            "Twilio must reach the app over HTTPS, so a public URL is required.",
            "cloudflared gives a free quick tunnel with no account and is tried first; ngrok is the fallback.",
            "Manual override is supported via PUBLIC_BASE_URL for anyone running their own tunnel.",
            "Why not deploy to a cloud host: an overnight hackathon values speed and zero-cost; a tunnel keeps the maker's credentials and data on their own machine.",
        ],
    )
    h2(doc, "7.5 FastAPI + Python")
    bullets(
        doc,
        [
            "Async-friendly for concurrent webhook calls and media downloads.",
            "Minimal boilerplate and automatic request parsing for form posts (Twilio) and multipart uploads (photos).",
            "Same language as the OpenAI SDK and the rest of the AI tooling, so a single file can contain the whole system.",
        ],
    )
    h2(doc, "7.6 Single-file app (app.py)")
    bullets(
        doc,
        [
            "The challenge is an overnight prototype: one command, no setup, no missing files.",
            "app.py auto-installs dependencies, embeds the backend and both UIs, starts the server, opens the browser, and (with --live) starts the tunnel.",
            "A modular backend/frontend version also exists for readability and extension.",
        ],
    )
    h2(doc, "7.7 In-memory sessions (trade-off)")
    bullets(
        doc,
        [
            "Why: fast to build, no database setup, perfect for a demo; sessions are keyed by phone number so multiple testers do not collide.",
            "Trade-off: sessions are lost on restart. A database (for example SQLite/Postgres) is the first production upgrade.",
        ],
    )
    h2(doc, "7.8 Shareable buyer link instead of a WhatsApp buyer flow")
    bullets(
        doc,
        [
            "The buyer gets a rich, visual page (photo, provenance) with no install and no opt-in.",
            "It keeps the maker as owner of the WhatsApp sender and avoids messaging-consent issues for buyers.",
            "It can be forwarded to anyone, which makes user testing with a friend trivial.",
        ],
    )
    h2(doc, "7.9 Voice notes with Whisper")
    bullets(
        doc,
        [
            "Directly addresses 'limited English, limited time': speaking is easier than typing.",
            "Whisper handles natural speech and multiple languages, enabling future non-English makers.",
        ],
    )
    h2(doc, "7.10 Why closed care vocabulary and explicit cultural fallback")
    bullets(
        doc,
        [
            "Care is a safety and satisfaction issue; an invented 'machine washable' causes ruined products and returns.",
            "Cultural meaning is a dignity issue; inventing heritage is worse than saying nothing.",
            "Both therefore use conservative, standardized language when the maker has not confirmed.",
        ],
    )

    # 8. Flows
    h1(doc, "8. User Flows")
    h2(doc, "8.1 Maker flow (WhatsApp)")
    numbered(
        doc,
        [
            "Maker sends a message (or voice note) to the WhatsApp number.",
            "Bot asks up to seven short questions with quick-reply options.",
            "Maker can send a photo; it is stored and shown on the buyer page.",
            "Each answer is echoed back for confirmation before it is recorded.",
            "Bot presents the guarded listing and lists any blocked claims.",
            "Maker publishes; bot returns the buyer link.",
        ],
    )
    h2(doc, "8.2 Buyer flow")
    numbered(
        doc,
        [
            "Buyer opens the shared link.",
            "Sees the product photo and the listing sections (story, materials, care, production, variations, cultural note, exact piece).",
            "Sees a trust banner and a provenance panel of the confirmed facts used.",
            "Sees a note when unconfirmed claims were deliberately blocked.",
        ],
    )
    h2(doc, "8.3 Owner and tester split")
    bullets(
        doc,
        [
            "Owner (you): holds the Twilio account and runs the app; tests the maker flow with photos and voice.",
            "Tester (friend): only needs the buyer link to act as the customer; no WhatsApp opt-in required.",
            "Multiple phones can use one sender; sessions are separated by phone number.",
        ],
    )

    # 9. API
    h1(doc, "9. API Reference")
    table(
        doc,
        ["Method & path", "Purpose"],
        [
            ("GET /", "Maker simulator UI (WhatsApp-style)."),
            ("GET /buyer/{id}", "Shareable buyer page."),
            ("GET /api/health", "Liveness and whether OpenAI/Twilio are configured."),
            ("POST /api/chat", "Simulator chat turn {session_id, message}."),
            ("POST /api/voice", "Upload a voice note; transcribes and continues the flow."),
            ("POST /api/photo", "Upload a product photo; stored and shown on the buyer page."),
            ("POST /api/reset", "Clear a session."),
            ("GET /api/listing/{id}", "Full listing: safe fields, draft, removed claims, provenance facts, photo URL."),
            ("GET /media/{id}", "Serves the stored product photo."),
            ("POST /webhook/twilio", "Twilio inbound webhook; returns TwiML replies."),
        ],
    )

    # 10. Security
    h1(doc, "10. Security, Privacy and Trust")
    bullets(
        doc,
        [
            "Secrets live only in .env, which is gitignored; credentials were never committed.",
            "Twilio credentials are used read-only to download incoming media.",
            "Product photos are held in memory and served only through the listing media route.",
            "No buyer-facing claim is published without a matching confirmed fact.",
            "Future hardening: validate the Twilio X-Twilio-Signature header, add authentication to maker/admin routes, and encrypt media at rest.",
        ],
    )

    # 11. Testing
    h1(doc, "11. Testing and Validation")
    bullets(
        doc,
        [
            "Offline smoke test (scripts/smoke_test.py and the single-file equivalent) covers interview, guard blocking, publish, buyer page, restart, retry, non-answer handling, voice, reset, invalid input, and the Twilio webhook - 28 checks, all passing.",
            "Command-line demo (scripts/demo.py) prints a full interview and the blocked claims.",
            "Verified live: public health ok, public webhook returns valid TwiML.",
            "Guard validation: on a maker who confirmed 'hand wash only', the guard blocked the invented 'It is also machine washable', the invented 'Ships within two days', and the invented 'symbolises prosperity and good fortune', while keeping the confirmed facts.",
        ],
    )

    # 12. Limitations
    h1(doc, "12. Limitations and Future Work")
    table(
        doc,
        ["Limitation", "Future work"],
        [
            ("In-memory sessions", "Persist sessions and the ledger in SQLite/Postgres."),
            ("Offline mock uses keyword heuristics", "Always run with a real model; add evaluation harness for the guard."),
            ("Trial limitations (verified numbers, ephemeral tunnel URL)", "Upgrade to a registered WhatsApp sender and a stable deployed URL."),
            ("English-first", "Extend voice + generation to the maker's native language."),
            ("Single product photo", "Support multiple photos and per-piece galleries."),
            ("No webhook signature validation", "Add X-Twilio-Signature verification and admin auth."),
            ("No analytics", "Track buyer questions and blocked-claim trends to improve listings."),
        ],
    )

    # 13. Demo
    h1(doc, "13. Demo Script")
    numbered(
        doc,
        [
            "Run `python app.py --live` (or `python app.py` for the offline simulator).",
            "Join the Twilio WhatsApp trial sender with the join code from your phone.",
            "Paste the printed tunnel URL + /webhook/twilio into the Twilio inbound webhook setting.",
            "Message the number: answer the questions, send a photo and a voice note.",
            "Observe the draft listing and the claims the guard blocked.",
            "Reply Publish and forward the buyer link to a tester.",
            "Open the buyer page to show the photo, sections, and the provenance panel.",
        ],
    )

    # 14. Glossary
    h1(doc, "14. Glossary")
    table(
        doc,
        ["Term", "Meaning"],
        [
            ("Claim Ledger", "The store of facts the maker explicitly confirmed; the only source of truth."),
            ("Atomic claim", "The smallest meaningful statement extracted from the listing draft."),
            ("Guard", "The audit-and-repair step that blocks or fixes unsupported claims."),
            ("Provenance", "The ability to trace a published sentence back to a confirmed fact."),
            ("TwiML", "Twilio Markup Language; the XML used to reply to an inbound WhatsApp message."),
            ("Tunnel", "A public HTTPS URL that forwards to the locally running app."),
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
